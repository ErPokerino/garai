"""Il motore non deve dipendere dal template Abstract: un template qualsiasi (qui costruito al volo, una sola slide,
slot e campi diversi, spec con errori tipici di una proposta LLM) deve produrre una presentazione."""
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Cm, Pt

from app.pipeline.fit import budget_feedback, fit_deck, trim_for_issue
from app.render.estimate import EstimateRenderer
from app.schemas import ExperienceBlock, PersonContent, TemplateSpec, WriterOutput
from app.schemas.content import ExtraField
from app.pipeline.tailor import assemble_person, normalize_writer_output
from app.template.budget import compute_budgets, template_fields
from app.template.deck import build_deck, make_meter
from app.template.normalize import normalize_spec


def _box(slide, x, y, w, h, text, size=14, bold=False):
    tb = slide.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
    tf = tb.text_frame
    for i, line in enumerate(text if isinstance(text, list) else [text]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size)
        r.font.bold = bold
    return tb.shape_id


@pytest.fixture()
def template(tmp_path) -> tuple[Path, dict]:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Cm(33.87), Cm(19.05)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    ids = {
        "name": _box(slide, 1, 1, 20, 1.2, "CANDIDATO", 28, True),
        "role": _box(slide, 1, 2.4, 20, 1, "Ruolo richiesto", 18),
        "info": _box(slide, 1, 4, 14, 2.5, ["Lingue:", "Disponibilità:"], 12),
        "skills_title": _box(slide, 1, 7, 14, 1, "Competenze", 16, True),
        "skills": _box(slide, 1, 8.2, 14, 0.8, ".", 12),
        "exp": _box(slide, 17, 4, 15, 0.8, ".", 12),
    }
    path = tmp_path / "gara.pptx"
    prs.save(str(path))
    return path, ids


def _proposed_spec(path: Path, ids: dict) -> TemplateSpec:
    """Spec come lo proporrebbe un LLM: alias di campo, shape inesistente, segnaposto da rimuovere invece che da riempire."""
    return TemplateSpec.model_validate({
        "name": "Gara", "template_file": str(path), "slides": [{"key": "cv", "index": 0}, {"key": "extra", "index": 7}],
        "labels": {"skills": {"it": "Competenze"}},
        "slots": [
            {"id": "name", "slide": "cv", "kind": "text", "field": "candidate_name", "shape_ids": [ids["name"]]},
            {"id": "role", "slide": "cv", "kind": "text", "field": "target_role", "shape_ids": [ids["role"], 999]},
            {"id": "info", "slide": "cv", "kind": "label_lines", "shape_ids": [ids["info"]],
             "options": {"lines": [{"label_key": "lang", "field": "languages"}, {"label_key": "avail", "field": "availability"}],
                         "separator": ": "}},
            {"id": "skills_title", "slide": "cv", "kind": "static_label", "shape_ids": [ids["skills_title"]],
             "options": {"label_key": "skills"}},
            {"id": "skills", "slide": "cv", "kind": "bullets", "field": "competences", "remove_shape_ids": [ids["skills"]],
             "box": {"x": 1, "y": 8.2, "w": 14, "h": 8}},
            {"id": "exp", "slide": "cv", "kind": "experience_columns", "field": "experiences", "shape_ids": [ids["exp"]]},
            {"id": "ghost", "slide": "nowhere", "kind": "text", "field": "email", "shape_ids": [1]},
        ],
    })


def _person(**kw) -> PersonContent:
    base = dict(
        source_file="cv.docx", profile_id="p1", profile_name="Project Manager", full_name="Giulia Verdi",
        skills=["Project management", "Agile Scrum", "SAP"],
        experiences=[ExperienceBlock(title="Energy - Portale", role="PM", period="2021 - oggi", bullets=["Guida del team", "Budget"]),
                     ExperienceBlock(title="Retail - E-commerce", role="PM", bullets=["Rilasci"])],
        extra={"languages": "Italiano, Inglese (C1)", "availability": "Immediata"},
    )
    base.update(kw)
    return PersonContent(**base)


def test_normalize_fixes_llm_proposal(template):
    path, ids = template
    spec, notes = normalize_spec(_proposed_spec(path, ids))
    assert [r.key for r in spec.slides] == ["cv"]  # slide inesistente scartata
    by_id = {s.id: s for s in spec.slots}
    assert "ghost" not in by_id
    assert by_id["name"].field == "full_name" and by_id["role"].field == "profile_name"
    assert by_id["role"].shape_ids == [ids["role"]]  # shape inesistente rimossa
    assert by_id["skills"].field == "skills" and by_id["skills"].shape_ids == [ids["skills"]]  # riempie il segnaposto
    assert by_id["exp"].kind == "experience_columns"
    assert {f.key for f in spec.fields} == {"languages", "availability"}
    assert spec.labels["lang"]["en"] == "Lingue"  # etichetta presa dal template
    assert 0 < spec.bottom_limit_cm <= 19.05


def test_budgets_follow_the_template(template):
    path, ids = template
    spec, _ = normalize_spec(_proposed_spec(path, ids))
    b = compute_budgets(spec, make_meter(spec))
    assert not b.shows("summary") and b.summary_max_chars == 0
    assert b.shows("skills") and b.skills_max_items >= 3
    assert b.exp_max_blocks >= 1
    assert set(b.custom_fields) == {"languages", "availability"}
    assert template_fields(spec)[:2] == ["full_name", "profile_name"]
    # nessun feedback sui campi che il template non stampa (es. sintesi lunga)
    assert not [f for f in budget_feedback(_person(summary="x" * 5000), b) if "summary" in f]


def test_deck_on_any_template(template, tmp_path):
    path, ids = template
    spec, _ = normalize_spec(_proposed_spec(path, ids))
    res = build_deck(spec, [_person(), _person(full_name="Marco Bianchi", source_file="b.docx")], {"skills": "Competenze"}, "it", tmp_path / "out.pptx")
    prs = Presentation(str(res.path))
    assert len(prs.slides) == 2
    text = " | ".join(sh.text_frame.text for sh in prs.slides[0].shapes if sh.has_text_frame)
    for expected in ("Giulia Verdi", "Project Manager", "Lingue: Italiano, Inglese (C1)", "Disponibilità: Immediata", "Agile Scrum", "Energy - Portale"):
        assert expected in text
    assert "CANDIDATO" not in text and "Ruolo richiesto" not in text


def test_fit_loop_and_trim_on_any_template(template, tmp_path):
    path, ids = template
    spec, _ = normalize_spec(_proposed_spec(path, ids))
    renderer = EstimateRenderer(make_meter(spec))
    many = [f"Competenza numero {i} con una descrizione decisamente lunga" for i in range(40)]
    fit = fit_deck(spec, [_person(skills=many)], {}, "it", tmp_path / "fit.pptx", renderer, rewriter=None)
    assert not fit.remaining  # ridotto finche' entra
    assert any("skills" in n for n in fit.notes[0])
    c = _person(extra={"languages": "Italiano. Inglese. Francese.", "availability": "Subito"})
    assert trim_for_issue(c, spec, "info", "", 0, 1, 10)  # label_lines: accorcia il valore piu' lungo


def test_writer_custom_fields_and_overrides(template):
    path, ids = template
    spec, _ = normalize_spec(_proposed_spec(path, ids))
    b = compute_budgets(spec, make_meter(spec))
    w = WriterOutput(current_role="Dev", summary="Non stampata", skills=["A"], experiences=[],
                     extra_fields=[ExtraField(key="languages", text="Italiano"), ExtraField(key="ignota", text="x")])
    w = normalize_writer_output(w, b)
    assert [f.key for f in w.extra_fields] == ["languages"]
    from app.schemas import CVCanonical, ProfileSpec

    cv = CVCanonical(source_file="cv.docx", full_name="Ada Rossi", headline="Developer")
    prof = ProfileSpec(id="p1", name="PM")
    c = assemble_person(cv, prof, w, "it", custom_types={"languages": "text"}, role_override="Lead Developer")
    assert c.extra == {"languages": "Italiano"} and c.current_role == "Lead Developer"
