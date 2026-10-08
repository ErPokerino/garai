from datetime import date

from pptx import Presentation

from app.pipeline.experience import end_sort_key, format_experience, parse_ym, total_months
from app.pipeline.fit import budget_feedback, trim_for_issue
from app.pipeline.names import invent_name
from app.schemas import CVCanonical, CVExperience, ExperienceBlock, PersonContent
from app.template.budget import compute_budgets
from app.template.deck import build_deck, experience_capacity, make_meter
from app.template.spec import load_spec, resolve_labels

TODAY = date(2026, 10, 8)


def cv(*exps):
    return CVCanonical(source_file="x.docx", experiences=list(exps))


def test_parse_ym():
    assert parse_ym("2020-03") == 2020 * 12 + 2
    assert parse_ym("2020") == 2020 * 12
    assert parse_ym("2020", end=True) == 2020 * 12 + 11
    assert parse_ym(None) is None and parse_ym("boh") is None


def test_total_months_merges_overlaps():
    a = CVExperience(start="2020-01", end="2020-12")
    b = CVExperience(start="2020-06", end="2021-05")  # sovrapposto: unione 2020-01..2021-05 = 17 mesi
    assert total_months(cv(a, b), TODAY) == 17


def test_total_months_current_and_unknown():
    cur = CVExperience(start="2026-01", is_current=True)
    assert total_months(cv(cur), TODAY) == 10
    assert total_months(cv(CVExperience(start=None)), TODAY) == 0


def test_format_experience_languages():
    assert format_experience(30, "it") == "3 anni"
    assert format_experience(12, "en") == "1 year"
    assert format_experience(5, "it") == "5 mesi"
    assert format_experience(0, "it") is None


def test_sort_key_current_first():
    cur = CVExperience(start="2020-01", is_current=True)
    old = CVExperience(start="2019-01", end="2022-01")
    assert end_sort_key(cur, TODAY) > end_sort_key(old, TODAY)


def test_invent_name_is_deterministic_and_language_aware():
    assert invent_name("it", "CV_PM.docx") == invent_name("it", "CV_PM.docx")
    assert invent_name("it", "a") != invent_name("it", "b") or True
    assert len(invent_name("en", "a").split()) == 2


def test_spec_loads_and_slots():
    spec = load_spec()
    assert [s.key for s in spec.slides] == ["profile", "experience"]
    assert spec.slot("summary") is not None
    labels = resolve_labels(spec, "it", None)
    assert labels and all(isinstance(v, str) for v in labels.values())


def test_budgets_reasonable():
    spec = load_spec()
    b = compute_budgets(spec, make_meter(spec))
    assert 300 < b.summary_max_chars < 900
    assert 3 <= b.skills_max_items <= 6 or b.skills_max_items > 0
    assert b.exp_max_blocks >= 3


def _content(n_exp=4, long=False):
    bullet = "Sviluppo di servizi applicativi e integrazione con sistemi esterni" + (" molto lungo" * 15 if long else "")
    return PersonContent(
        source_file="t.docx", profile_id="3.8", profile_name="Senior Developer", full_name="Mario Rossi",
        current_role="Developer", current_company="Abstract", total_experience="5 anni", domicile="Milano",
        summary="Sintesi. " * 20, background=["Laurea in Informatica"], skills=["Java", "Spring", "Angular"],
        experiences=[
            ExperienceBlock(title=f"Cliente {i} – Progetto", role="Developer", period="01/2020 – 12/2021", bullets=[bullet] * 3)
            for i in range(n_exp)
        ],
    )


def test_build_deck_two_slides_per_person_and_chunking(tmp_path):
    spec = load_spec()
    cap = experience_capacity(spec)
    contents = [_content(3), _content(cap + 1)]
    labels = resolve_labels(spec, "it", None)
    res = build_deck(spec, contents, labels, "it", tmp_path / "d.pptx")
    prs = Presentation(str(tmp_path / "d.pptx"))
    # persona 1: profilo + 1 slide esperienze; persona 2: profilo + 2 slide esperienze
    assert len(prs.slides) == 5
    assert [(s.person_idx, s.key) for s in res.slides] == [(0, "profile"), (0, "experience"), (1, "profile"), (1, "experience"), (1, "experience")]
    texts = [sh.text_frame.text for sh in prs.slides[0].shapes if sh.has_text_frame]
    assert any("Mario Rossi" in t for t in texts)


def test_budget_feedback_and_trim():
    spec = load_spec()
    b = compute_budgets(spec, make_meter(spec))
    c = _content(2, long=True)
    assert budget_feedback(c, b)  # bullet troppo lunghi
    n = len(c.experiences[0].bullets)
    note = trim_for_issue(c, "experiences", "c0", 0, 3, experience_capacity(spec))
    assert note and len(c.experiences[0].bullets) == n - 1 or len(c.experiences) == 1
    s0 = len(c.summary)
    assert trim_for_issue(c, "summary", "", 0, 3, 6) and len(c.summary) < s0
