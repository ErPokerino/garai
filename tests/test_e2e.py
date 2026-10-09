"""End-to-end offline con le risposte LLM prefabbricate (samples/fixtures): bando + 6 CV -> PPTX."""
from pathlib import Path

import pytest
from pptx import Presentation

from app.llm.client import build_client
from app.pipeline.run import Pipeline

S = Path(__file__).resolve().parents[1] / "samples"
CVS = ["CV_DEV.docx", "CV_PM.docx", "CV_SA02.pdf", "CVDEV03.docx", "CVSA01.docx", "CVSA03.docx"]

pytestmark = pytest.mark.skipif(not (S / "fixtures").exists(), reason="fixtures mancanti")


@pytest.fixture(scope="module")
def run(tmp_path_factory):
    out = tmp_path_factory.mktemp("e2e")
    pipe = Pipeline(build_client(fixtures_dir=S / "fixtures", use_cache=False))
    _, bando = pipe.read_bando(next(S.glob("Annex B ITA*.docx")))
    items = pipe.parse_cvs([S / n for n in CVS], out)
    assigns = pipe.assign(items, bando)
    res = pipe.generate(bando, items, assigns, out, visual_critic=False)
    return bando, items, assigns, res


def test_bando_language_and_profiles(run):
    bando, *_ = run
    assert bando.language == "it"
    assert len(bando.profiles) >= 10


def test_one_pptx_with_all_cvs(run):
    _, items, _, res = run
    prs = Presentation(str(res.pptx_path))
    assert len(prs.slides) >= 2 * len(items)
    assert len(res.report.people) == len(CVS)


def test_no_residual_overflow(run):
    *_, res = run
    assert all(not p.fit_issues for p in res.report.people)


def test_missing_names_are_left_empty_and_flagged(run):
    """Il nome non si inventa: se il CV non lo riporta resta vuoto ed e' segnalato."""
    *_, res = run
    assert any(p.content.name_is_placeholder for p in res.report.people)
    for p in res.report.people:
        if p.content.name_is_placeholder:
            assert p.content.full_name == ""
            assert any("lasciato vuoto" in w for w in p.content.warnings)
        else:
            assert p.content.full_name


def test_all_shapes_inside_slide(run):
    *_, res = run
    prs = Presentation(str(res.pptx_path))
    W, H = prs.slide_width, prs.slide_height
    for sl in prs.slides:
        for sh in sl.shapes:
            if sh.width is None or sh.left is None:
                continue
            assert sh.left >= -10_000 and sh.top >= -10_000
            assert sh.left + sh.width <= W + 10_000 and sh.top + sh.height <= H + 10_000


def test_no_template_placeholders_left(run):
    *_, res = run
    prs = Presentation(str(res.pptx_path))
    for sl in prs.slides:
        for sh in sl.shapes:
            if sh.has_text_frame:
                assert "Nome Cognome" not in sh.text_frame.text
