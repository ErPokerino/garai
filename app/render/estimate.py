"""Misura del testo stimata con Pillow (fallback quando PowerPoint non e' disponibile) e renderer LibreOffice."""
from __future__ import annotations

import os
import shutil
import subprocess
import tempfile
from pathlib import Path

import pymupdf
from pptx import Presentation
from pptx.util import Emu

from app.template.measure import CM_TO_PT, TextMeter
from app.template.xmlutil import iter_shapes, paragraphs, qa, txbody

from .base import ShapeMeasure

EMU_PER_PT = 12700
DEFAULT_PT = 10.0
# La stima va a capo un po' prima di PowerPoint (larghezza utile ridotta): meglio una riga in meno che un testo che esce.
WIDTH_MARGIN = float(os.environ.get("ESTIMATE_WIDTH_MARGIN", "0.04"))


def _paragraph_lines(meter: TextMeter, p, width_pt: float) -> tuple[int, float]:
    runs = p.findall(qa("r"))
    text = "".join((r.find(qa("t")).text or "") for r in runs if r.find(qa("t")) is not None)
    size = DEFAULT_PT
    bold = False
    for r in runs:
        rpr = r.find(qa("rPr"))
        if rpr is not None:
            if rpr.get("sz"):
                size = int(rpr.get("sz")) / 100
            if rpr.get("b") == "1":
                bold = True
            break
    ppr = p.find(qa("pPr"))
    marl = int(ppr.get("marL", "0")) / EMU_PER_PT if ppr is not None else 0.0
    if not text.strip():
        return 1, size
    lines = meter.wrap(text, max(width_pt - marl, 10), size, bold)
    return len(lines), size


class EstimateMeasurer:
    """Calcola ShapeMeasure leggendo le shape marcate `slot:` e simulando il wrapping."""

    def __init__(self, meter: TextMeter):
        self.meter = meter

    def measure(self, pptx_path: Path) -> list[ShapeMeasure]:
        prs = Presentation(str(pptx_path))
        out: list[ShapeMeasure] = []
        for si, slide in enumerate(prs.slides, 1):
            for sh in iter_shapes(slide.shapes):
                if not sh.name.startswith("slot:") or "|chk=" not in sh.name or not sh.has_text_frame:
                    continue
                tf = sh.text_frame
                if not tf.text.strip():
                    continue
                body = txbody(sh)
                bpr = body.find(qa("bodyPr"))
                l_ins = int(bpr.get("lIns", 91440)) if bpr is not None else 91440
                r_ins = int(bpr.get("rIns", 91440)) if bpr is not None else 91440
                t_ins = int(bpr.get("tIns", 45720)) if bpr is not None else 45720
                width_pt = (sh.width - l_ins - r_ins) / EMU_PER_PT * (1 - WIDTH_MARGIN)
                total_lines, total_h = 0, 0.0
                paras = list(paragraphs(body))
                while len(paras) > 1 and not "".join(t.text or "" for t in paras[-1].iter(qa("t"))).strip():
                    paras.pop()  # paragrafi vuoti finali del template: non contano come righe di testo
                for p in paras:
                    n, size = _paragraph_lines(self.meter, p, width_pt)
                    total_lines += n
                    total_h += n * self.meter.line_height_pt(size)
                out.append(
                    ShapeMeasure(
                        slide_no=si,
                        name=sh.name,
                        shape_top_pt=sh.top / EMU_PER_PT,
                        shape_height_pt=sh.height / EMU_PER_PT,
                        text_top_pt=(sh.top + t_ins) / EMU_PER_PT,
                        text_height_pt=total_h,
                        n_lines=total_lines,
                    )
                )
        return out


class LibreOfficeRenderer:
    """Rendering con LibreOffice headless (PDF -> PNG); misure stimate con Pillow."""

    name = "libreoffice"

    def __init__(self, meter: TextMeter):
        self.exe = shutil.which("soffice") or shutil.which("libreoffice")
        if not self.exe:
            raise RuntimeError("LibreOffice (soffice) non trovato nel PATH")
        self._est = EstimateMeasurer(meter)

    def render(self, pptx_path: Path, out_dir: Path, width_px: int = 1600) -> list[Path]:
        out_dir = Path(out_dir)
        if out_dir.exists():
            shutil.rmtree(out_dir, ignore_errors=True)
        out_dir.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory() as td:
            subprocess.run(
                [self.exe, "--headless", "--convert-to", "pdf", "--outdir", td, str(pptx_path)],
                check=True,
                capture_output=True,
                timeout=240,
            )
            pdf = Path(td) / (Path(pptx_path).stem + ".pdf")
            doc = pymupdf.open(str(pdf))
            for i, page in enumerate(doc, 1):
                zoom = width_px / page.rect.width
                page.get_pixmap(matrix=pymupdf.Matrix(zoom, zoom)).save(str(out_dir / f"slide_{i:03d}.png"))
        return sorted(out_dir.glob("slide_*.png"))

    def measure(self, pptx_path: Path) -> list[ShapeMeasure]:
        return self._est.measure(pptx_path)

    def close(self) -> None:
        pass


class EstimateRenderer:
    """Nessun rendering (solo stima): usato quando non c'e' ne' PowerPoint ne' LibreOffice."""

    name = "estimate"

    def __init__(self, meter: TextMeter):
        self._est = EstimateMeasurer(meter)

    def render(self, pptx_path: Path, out_dir: Path, width_px: int = 1600) -> list[Path]:
        return []

    def measure(self, pptx_path: Path) -> list[ShapeMeasure]:
        return self._est.measure(pptx_path)

    def close(self) -> None:
        pass


def get_renderer(meter: TextMeter):
    """PowerPoint (se presente) > LibreOffice > stima Pillow. Forzabile con RENDERER=powerpoint|libreoffice|estimate."""
    from .powerpoint import PowerPointRenderer, available

    forced = os.environ.get("RENDERER", "").lower()
    if forced == "estimate":
        return EstimateRenderer(meter)
    if forced == "libreoffice":
        return LibreOfficeRenderer(meter)
    if available() and forced in ("", "powerpoint"):
        try:
            return PowerPointRenderer()
        except Exception:
            pass
    try:
        return LibreOfficeRenderer(meter)
    except Exception:
        return EstimateRenderer(meter)


_ = Emu, CM_TO_PT  # re-export per comodita'
