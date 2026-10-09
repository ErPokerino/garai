"""Validazione e correzione deterministica di un Template Spec (tipicamente proposto da un LLM).

L'LLM descrive il template; questo modulo garantisce che la descrizione sia eseguibile: slide e shape esistenti,
campi noti o dichiarati, etichette presenti, aree dentro la slide, font del template per le misure.
Cio' che non si puo' correggere viene scartato con una nota, invece di far fallire la generazione.
"""
from __future__ import annotations

import re
from pathlib import Path

from pptx import Presentation

from app.schemas import STANDARD_FIELDS, Box, CustomField, SlideRef, Slot, TemplateSpec

from .fill import para_size_pt
from .fonts import dominant_font, resolve_font_files
from .geometry import auto_bottom_limit, cm, slide_size_cm
from .xmlutil import paragraphs, qa, shape_by_id, txbody

FIELD_ALIASES = {
    "name": "full_name", "candidate_name": "full_name", "fullname": "full_name",
    "profile": "profile_name", "target_role": "profile_name", "requested_role": "profile_name", "position": "profile_name",
    "role": "current_role", "job_title": "current_role", "title": "current_role",
    "company": "current_company", "employer": "current_company",
    "location": "domicile", "city": "domicile", "residence": "domicile",
    "years_of_experience": "total_experience", "experience_years": "total_experience", "total_years": "total_experience",
    "experience": "experiences", "work_experience": "experiences", "professional_experience": "experiences",
    "profile_summary": "summary", "about": "summary",
    "competences": "skills", "competencies": "skills", "technical_skills": "skills",
    "mobile": "phone", "telephone": "phone", "mail": "email",
}
NEEDS_SHAPE = ("static_label", "text", "label_lines", "pills")


def _field(name: str | None) -> str | None:
    if not name:
        return None
    key = re.sub(r"[^a-z0-9_]", "_", name.strip().lower()).strip("_")
    return FIELD_ALIASES.get(key, key)


def _shape_text(sh) -> str:
    try:
        return " ".join("".join(t.text or "" for t in p.iter(qa("t"))) for p in paragraphs(txbody(sh))).strip()
    except ValueError:
        return ""


def _overlaps(sh, box: Box) -> bool:
    x, y, w, h = cm(sh.left), cm(sh.top), cm(sh.width), cm(sh.height)
    return min(x + w, box.x + box.w) > max(x, box.x) and min(y + max(h, 0.5), box.y + box.h) > max(y, box.y)


def normalize_spec(spec: TemplateSpec, pptx_path: str | Path | None = None) -> tuple[TemplateSpec, list[str]]:
    spec = spec.model_copy(deep=True)
    if pptx_path is not None:
        spec.template_file = str(pptx_path)
    prs = Presentation(str(spec.template_file))
    notes: list[str] = []
    W, H = slide_size_cm(prs)

    # --- slide del set persona
    refs: list[SlideRef] = []
    for r in spec.slides:
        if 0 <= r.index < len(prs.slides) and r.key not in {x.key for x in refs}:
            refs.append(r)
    if not refs:
        refs = [SlideRef(key=spec.slots[0].slide if spec.slots else "profile", index=0)]
        notes.append("Nessuna slide valida indicata: uso la prima del template.")
    spec.slides = refs
    keys = {r.key for r in refs}

    # --- limite inferiore (piede di pagina)
    auto = min(auto_bottom_limit(prs, prs.slides[r.index]) for r in refs)
    spec.bottom_limit_cm = auto if not (0.5 * H <= spec.bottom_limit_cm <= H) else min(spec.bottom_limit_cm, auto)
    spec.line_height_factor = min(1.35, max(1.1, spec.line_height_factor or 1.2))

    # --- slot
    out: list[Slot] = []
    used_ids: set[str] = set()
    for slot in spec.slots:
        if slot.slide not in keys:
            if len(keys) == 1:
                slot.slide = next(iter(keys))
            else:
                notes.append(f"Slot '{slot.id}' su una slide inesistente: scartato.")
                continue
        slide = prs.slides[next(r.index for r in refs if r.key == slot.slide)]
        exists = lambda i: shape_by_id(slide, i) is not None  # noqa: E731
        slot.field = _field(slot.field)
        slot.shape_ids = [i for i in slot.shape_ids if exists(i)]
        slot.remove_shape_ids = [i for i in slot.remove_shape_ids if exists(i) and i not in slot.shape_ids]
        slot.remove_shape_ids_if_empty = [i for i in slot.remove_shape_ids_if_empty if exists(i)]
        if slot.kind == "label_lines":
            lines = []
            for ln in slot.options.get("lines") or []:
                if isinstance(ln, dict) and ln.get("field"):
                    lines.append({**ln, "field": _field(ln["field"])})
            slot.options["lines"] = lines
            if not lines:
                notes.append(f"Slot '{slot.id}' senza righe: scartato.")
                continue

        if slot.kind == "experience_columns" and (slot.field != "experiences" or not slot.shape_ids):
            slot.kind = "bullets"
        if slot.kind == "pills":
            opts = slot.options
            if not (slot.shape_ids and opts.get("columns") and opts.get("first_row_y") is not None
                    and opts.get("row_pitch") and opts.get("pill_h")):
                slot.kind = "bullets"  # pillole descritte in modo incompleto: elenco nell'area della pillola
                notes.append(f"Slot '{slot.id}': pillole incomplete, uso un elenco.")
        if slot.kind in ("bullets", "free_text") and not slot.shape_ids:
            # il segnaposto d'esempio che l'LLM voleva rimuovere e' la shape giusta da riempire (ne conserva lo stile)
            cand = [i for i in slot.remove_shape_ids if (sh := shape_by_id(slide, i)) is not None and sh.has_text_frame
                    and (slot.box is None or _overlaps(sh, slot.box))]
            if cand:
                slot.shape_ids = [cand[0]]
                slot.remove_shape_ids = [i for i in slot.remove_shape_ids if i != cand[0]]
        if slot.kind in NEEDS_SHAPE and not slot.shape_ids:
            notes.append(f"Slot '{slot.id}' senza shape valida: scartato.")
            continue
        if slot.kind in ("bullets", "free_text") and not slot.shape_ids and slot.box is None:
            notes.append(f"Slot '{slot.id}' senza area: scartato.")
            continue
        if slot.kind != "static_label" and not slot.field and slot.kind != "label_lines":
            notes.append(f"Slot '{slot.id}' senza campo: scartato.")
            continue
        if slot.box is not None:
            b = slot.box
            x, y = min(max(0.0, b.x), W - 1), min(max(0.0, b.y), H - 1)
            slot.box = Box(x=x, y=y, w=max(0.5, min(b.w, W - x)), h=max(0.3, min(b.h, H - y)))
        if slot.font_pt is None and slot.shape_ids:
            sh = shape_by_id(slide, slot.shape_ids[0])
            if sh is not None and sh.has_text_frame:
                ps = paragraphs(txbody(sh))
                slot.font_pt = para_size_pt(ps[0]) if ps else None
        sid = slot.id or slot.kind
        while sid in used_ids:
            sid += "_2"
        slot.id = sid
        used_ids.add(sid)
        out.append(slot)
    spec.slots = out

    # --- campi personalizzati dichiarati
    declared = {f.key: f for f in spec.fields}
    fields: list[CustomField] = []
    for slot in out:
        names = [ln["field"] for ln in slot.options.get("lines", [])] if slot.kind == "label_lines" else [slot.field]
        for n in names:
            if not n or n in STANDARD_FIELDS or n in {f.key for f in fields}:
                continue
            f = declared.get(n) or CustomField(
                key=n, description=n.replace("_", " ").capitalize(),
                type="list" if slot.kind in ("bullets", "pills") else "text",
            )
            fields.append(f)
    spec.fields = fields

    # --- etichette: ogni label_key usato deve esistere (si ricade sul testo del template)
    for slot in out:
        slide = prs.slides[next(r.index for r in refs if r.key == slot.slide)]
        if slot.kind == "static_label":
            key = slot.options.get("label_key") or slot.id
            slot.options["label_key"] = key
            if key not in spec.labels:
                text = _shape_text(shape_by_id(slide, slot.shape_ids[0]))
                if text:
                    spec.labels[key] = {"en": text}
        elif slot.kind == "label_lines":
            sh = shape_by_id(slide, slot.shape_ids[0])
            ps = paragraphs(txbody(sh))
            for i, ln in enumerate(slot.options["lines"]):
                key = ln.get("label_key")
                if key and key not in spec.labels and i < len(ps):
                    text = "".join(t.text or "" for t in ps[i].iter(qa("t"))).strip().rstrip(":").strip()
                    if text:
                        spec.labels[key] = {"en": text}

    # --- font del template per le misure stimate e per i testi creati da zero
    family = dominant_font(prs, [prs.slides[r.index] for r in refs])
    if family:
        files = resolve_font_files(family)
        if files:
            spec.font_regular, spec.font_bold = files
            spec.font_family = family
        else:
            notes.append(f"Font '{family}' non disponibile: misure stimate con il font predefinito.")
    return spec, notes
