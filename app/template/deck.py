"""Costruzione del deck finale: clona le slide del template per ogni persona e le riempie."""
from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from pptx import Presentation

from app.schemas import FitIssue, PersonContent, TemplateSpec

from .budget import compute_budgets
from .clone import clone_slide, delete_slides
from .fill import FillContext, extract_protos, fill_slide, slot_areas
from .measure import TextMeter
from .spec import template_path


@dataclass
class SlideInfo:
    person_idx: int
    key: str
    slide_no: int  # 1-based nel deck finale
    chunk: int = 0  # indice del gruppo di esperienze mostrato (slide esperienze multiple)


@dataclass
class DeckResult:
    path: Path
    slides: list[SlideInfo]
    issues: list[FitIssue]


def make_meter(spec: TemplateSpec) -> TextMeter:
    return TextMeter(spec.font_regular, spec.font_bold, spec.line_height_factor)


def experience_columns_slot(spec: TemplateSpec):
    """Slot a colonne di esperienze (se il template ne ha uno): e' l'unico che si ripete su piu' slide."""
    return next((s for s in spec.slots if s.kind == "experience_columns" and s.shape_ids), None)


def experience_capacity(spec: TemplateSpec) -> int:
    ex = experience_columns_slot(spec)
    if ex is None:
        return 10_000
    return len(ex.shape_ids) * int(ex.options.get("blocks_per_column", 2))


def build_deck(
    spec: TemplateSpec,
    contents: list[PersonContent],
    labels: dict[str, str],
    language: str,
    out_path: str | Path,
) -> DeckResult:
    meter = make_meter(spec)
    prs = Presentation(str(template_path(spec)))
    protos = extract_protos(prs, spec)
    areas = slot_areas(prs, spec)
    originals = {ref.key: prs.slides[ref.index] for ref in spec.slides}
    n_orig = len(prs.slides)
    cap = experience_capacity(spec)
    ex_cols = experience_columns_slot(spec)

    infos: list[SlideInfo] = []
    issues: list[FitIssue] = []
    slide_no = 0
    for pi, content in enumerate(contents):
        for ref in spec.slides:
            src = originals[ref.key]
            if ex_cols is not None and ex_cols.slide == ref.key:  # esperienze oltre la capienza: slide in piu'
                exps = content.experiences
                chunks = [exps[i : i + cap] for i in range(0, max(len(exps), 1), cap)] or [[]]
            else:
                chunks = [None]
            for chunk_idx, chunk in enumerate(chunks):
                slide_no += 1
                slide = clone_slide(prs, src)
                ctx = FillContext(
                    spec=spec,
                    labels=labels,
                    lang=language,
                    protos=protos,
                    areas=areas,
                    meter=meter,
                    content=content,
                    person_idx=pi,
                    slide_no=slide_no,
                    exp_chunk=chunk,
                )
                fill_slide(slide, ref.key, ctx)
                issues.extend(ctx.issues)
                infos.append(SlideInfo(pi, ref.key, slide_no, chunk_idx))

    # le slide originali del template restano in testa: vengono rimosse
    delete_slides(prs, list(range(n_orig)))
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))
    return DeckResult(path=out_path, slides=infos, issues=issues)


__all__ = ["build_deck", "DeckResult", "SlideInfo", "make_meter", "experience_capacity", "compute_budgets"]
