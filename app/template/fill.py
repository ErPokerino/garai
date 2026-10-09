"""Riempimento delle slide clonate dal template secondo il Template Spec.

Principio: non si genera mai una slide da zero. Si clonano le slide del template, si sostituisce il testo
conservando run e stili originali e si clonano i blocchi ripetibili (pillole, paragrafi, colonne di esperienza).
Ogni slot prende i propri prototipi (paragrafo, punto elenco, pillola) dalla propria shape del template: nessuno slot
dipende da un altro, quindi qualsiasi combinazione di slot funziona.
Le shape da controllare sul rendering reale vengono marcate nel nome (`slot:<id>:<persona>[:extra]|chk=<regola>`),
cosi' il controllo dell'overflow sa cosa verificare.
"""
from __future__ import annotations

import copy
from dataclasses import dataclass, field

from lxml import etree
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.presentation import Presentation
from pptx.slide import Slide
from pptx.util import Cm, Emu

from app.schemas import STANDARD_FIELDS, Box, ExperienceBlock, FitIssue, PersonContent, Slot, TemplateSpec

from .geometry import slide_of, slot_area
from .measure import CM_TO_PT, TextMeter
from .xmlutil import (
    A_NS,
    append_run,
    clear_paragraph_content,
    first_rpr,
    iter_shapes,
    make_run,
    max_shape_id,
    paragraph_from_proto,
    paragraphs,
    qa,
    remove_shape,
    set_cnv_id,
    set_shape_name,
    shape_by_id,
    style_rpr,
    txbody,
)

LANG_TAGS = {
    "it": "it-IT", "en": "en-US", "es": "es-ES", "fr": "fr-FR", "de": "de-DE",
    "pt": "pt-PT", "ro": "ro-RO",
}
INSET_CM = 0.1
PILL_INSET_CM = 0.3
BULLET_MAR_EMU = 228600  # rientro di un punto elenco aggiunto quando il template non ne ha uno
SUB_STEP_EMU = 342900  # rientro aggiuntivo dei sotto-punti
BULLET_KINDS = ("buChar", "buAutoNum", "buBlip")


def lang_tag(lang: str) -> str:
    return LANG_TAGS.get(lang, lang)


def shorten(text: str, max_chars: int) -> str:
    if len(text) <= max_chars:
        return text
    cut = text[: max_chars - 1].rsplit(" ", 1)[0].rstrip(" ,;:-")
    return cut + "…"


# ----------------------------------------------------------------------------- prototipi per slot


def para_size_pt(p: etree._Element | None) -> float | None:
    if p is None:
        return None
    for el in list(p.iter(qa("rPr"))) + list(p.iter(qa("endParaRPr"))) + list(p.iter(qa("defRPr"))):
        if el.get("sz"):
            return int(el.get("sz")) / 100
    return None


def _has_bullet(p: etree._Element) -> bool:
    ppr = p.find(qa("pPr"))
    return ppr is not None and any(ppr.find(qa(k)) is not None for k in BULLET_KINDS)


def _synth_paragraph() -> etree._Element:
    p = etree.Element(qa("p"))
    etree.SubElement(p, qa("pPr"))
    return p


def _with_bullet(p: etree._Element, char: str = "•", extra_mar: int = 0) -> etree._Element:
    """Copia del paragrafo con punto elenco (aggiunto se il template non lo prevede) e rientro eventuale."""
    p = copy.deepcopy(p)
    ppr = p.find(qa("pPr"))
    if ppr is None:
        ppr = etree.Element(qa("pPr"))
        p.insert(0, ppr)
    if not _has_bullet(p):
        for tag in ("buNone",):
            el = ppr.find(qa(tag))
            if el is not None:
                ppr.remove(el)
        ppr.set("marL", str(int(ppr.get("marL", BULLET_MAR_EMU))))
        ppr.set("indent", str(int(ppr.get("indent", -BULLET_MAR_EMU))))
        bu = etree.SubElement(ppr, qa("buChar"))
        bu.set("char", char)
    if extra_mar:
        ppr.set("marL", str(int(ppr.get("marL", 0)) + extra_mar))
        bu = ppr.find(qa("buChar"))
        if bu is not None:
            bu.set("char", char)
    return p


def _first_text_para(shape) -> etree._Element | None:
    try:
        ps = paragraphs(txbody(shape))
    except ValueError:
        return None
    for p in ps:
        if "".join(t.text or "" for t in p.iter(qa("t"))).strip():
            return p
    return ps[0] if ps else None


def _any_bullet_para(slides: list[Slide]) -> etree._Element | None:
    """Primo punto elenco presente nel template (prima nella slide dello slot, poi nelle altre)."""
    for slide in slides:
        for sh in iter_shapes(slide.shapes):
            if not sh.has_text_frame:
                continue
            for p in paragraphs(txbody(sh)):
                if _has_bullet(p):
                    return p
    return None


def extract_protos(prs: Presentation, spec: TemplateSpec) -> dict[str, dict[str, etree._Element]]:
    """Prototipi presi dalle slide originali del template, prima della clonazione, per ogni slot."""
    out: dict[str, dict[str, etree._Element]] = {}
    for slot in spec.slots:
        slide = slide_of(prs, spec, slot)
        sh = shape_by_id(slide, slot.shape_ids[0]) if slot.shape_ids else None
        own = _first_text_para(sh) if sh is not None and sh.has_text_frame else None
        pr: dict[str, etree._Element] = {}
        if slot.kind in ("bullets", "free_text"):
            others = [prs.slides[r.index] for r in spec.slides if r.key != slot.slide]
            base = own if own is not None else (_any_bullet_para([slide, *others]) if slot.kind == "bullets" else None)
            pr["para"] = copy.deepcopy(base) if base is not None else _synth_paragraph()
        elif slot.kind == "pills" and sh is not None:
            pr["pill"] = copy.deepcopy(sh._element)
        elif slot.kind == "experience_columns":
            col = shape_by_id(slide, int(slot.options.get("proto_column", slot.shape_ids[0] if slot.shape_ids else -1)))
            ps = paragraphs(txbody(col)) if col is not None and col.has_text_frame else []
            bullet = next((p for p in ps if _has_bullet(p)), None)
            plain = ps[0] if ps else _synth_paragraph()
            pr["title"] = copy.deepcopy(ps[0] if ps else plain)
            pr["role"] = copy.deepcopy(ps[1] if len(ps) > 1 else plain)
            pr["blank"] = copy.deepcopy(ps[2] if len(ps) > 2 else plain)
            pr["bullet"] = copy.deepcopy(bullet if bullet is not None else _with_bullet(plain))
        out[slot.id] = pr
    return out


@dataclass
class FillContext:
    spec: TemplateSpec
    labels: dict[str, str]
    lang: str
    protos: dict[str, dict[str, etree._Element]]
    areas: dict[str, Box | None]
    meter: TextMeter
    content: PersonContent
    person_idx: int
    slide_no: int = 0  # numero (1-based) della slide nel deck finale
    exp_chunk: list[ExperienceBlock] | None = None
    issues: list[FitIssue] = field(default_factory=list)

    @property
    def tag(self) -> str:
        return lang_tag(self.lang)

    def warn(self, slot: str, detail: str, severity: str = "warning") -> None:
        self.issues.append(FitIssue(slide=self.slide_no, slot=slot, detail=detail, severity=severity))


# ----------------------------------------------------------------------------- valori dei campi


def field_value(content: PersonContent, name: str | None):
    if not name:
        return None
    if name in STANDARD_FIELDS:
        return getattr(content, name, None)
    return content.extra.get(name)


def exp_header(e: ExperienceBlock) -> str:
    role = e.role + (f" ({e.period})" if e.period else "")
    return f"{e.title} – {role}" if e.title and role else (e.title or role)


def as_text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, list):
        if v and isinstance(v[0], ExperienceBlock):
            return "; ".join(exp_header(e) for e in v)
        return ", ".join(str(x) for x in v if str(x).strip())
    return str(v).strip()


def as_items(v) -> list:
    """Voci di un elenco: blocchi esperienza cosi' come sono, stringhe ripulite, testo diviso per righe."""
    if not v:
        return []
    if isinstance(v, list):
        out = []
        for x in v:
            if isinstance(x, ExperienceBlock):
                if x.title or x.role:
                    out.append(x)
            elif str(x).strip():
                out.append(str(x).strip())
        return out
    return [s.strip() for s in str(v).split("\n") if s.strip()]


# ----------------------------------------------------------------------------- helpers


def _name(slot_id: str, ctx: FillContext, extra: str = "", chk: str | None = None) -> str:
    n = f"slot:{slot_id}:{ctx.person_idx}" + (f":{extra}" if extra else "")
    return n + (f"|chk={chk}" if chk else "")


def _clean_end(p: etree._Element) -> None:
    end = p.find(qa("endParaRPr"))
    if end is not None and "err" in end.attrib:
        del end.attrib["err"]


def _set_single_text(shape, text: str, ctx: FillContext) -> None:
    tb = txbody(shape)
    ps = paragraphs(tb)
    p0 = _first_text_para(shape) if ps else None
    if p0 is None:
        p0 = etree.SubElement(tb, qa("p"))
    rpr = first_rpr(p0)
    if rpr is None:
        end = p0.find(qa("endParaRPr"))
        rpr = copy.deepcopy(end) if end is not None else None
        if rpr is not None:
            rpr.tag = qa("rPr")
    clear_paragraph_content(p0)
    append_run(p0, make_run(text, style_rpr(rpr, lang=ctx.tag)))
    _clean_end(p0)
    for p in paragraphs(tb):  # eventuali altri paragrafi d'esempio del template
        if p is not p0:
            tb.remove(p)


def _remove_ids(slide: Slide, ids: list[int]) -> None:
    for sid in ids:
        sh = shape_by_id(slide, sid)
        if sh is not None:
            remove_shape(sh)


def _no_autofit(shape) -> None:
    """Dimensione fissa: il testo deve stare nell'area, non far crescere la shape (il controllo e' sull'altezza)."""
    bpr = txbody(shape).find(qa("bodyPr"))
    if bpr is None:
        return
    for tag in ("spAutoFit", "normAutofit", "noAutofit"):
        el = bpr.find(qa(tag))
        if el is not None:
            bpr.remove(el)
    etree.SubElement(bpr, qa("noAutofit"))
    bpr.set("wrap", "square")


def _target_shape(slide: Slide, slot: Slot, ctx: FillContext, extra_name: str = ""):
    """Shape in cui scrivere: quella del template (ridimensionata all'area utile) oppure una nuova casella di testo."""
    area = ctx.areas.get(slot.id)
    sh = shape_by_id(slide, slot.shape_ids[0]) if slot.shape_ids else None
    if sh is not None and sh.has_text_frame:
        if area is not None:
            sh.left, sh.top, sh.width, sh.height = Cm(area.x), Cm(area.y), Cm(area.w), Cm(area.h)
        _no_autofit(sh)
        set_shape_name(sh, _name(slot.id, ctx, extra_name, chk="h"))
        return sh
    if area is None:
        ctx.warn(slot.id, "slot senza area o shape valida: ignorato")
        return None
    tb = slide.shapes.add_textbox(Cm(area.x), Cm(area.y), Cm(area.w), Cm(area.h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    set_shape_name(tb, _name(slot.id, ctx, extra_name, chk="h"))
    return tb


def _font(ctx: FillContext) -> str | None:
    return ctx.spec.font_family


def _spacing(p: etree._Element, space_after_pt: float) -> None:
    ppr = p.find(qa("pPr"))
    if ppr is None:
        ppr = etree.Element(qa("pPr"))
        p.insert(0, ppr)
    for tag in ("spcBef", "spcAft"):
        el = ppr.find(qa(tag))
        if el is not None:
            ppr.remove(el)
    if space_after_pt:
        spc = etree.Element(qa("spcAft"))
        etree.SubElement(spc, qa("spcPts")).set("val", str(int(space_after_pt * 100)))
        ln = ppr.find(qa("lnSpc"))  # ordine dello schema: lnSpc, spcBef, spcAft, bu*
        if ln is not None:
            ln.addnext(spc)
        else:
            ppr.insert(0, spc)


def _replace_paragraphs(shape, new_ps: list[etree._Element]) -> None:
    tb = txbody(shape)
    for p in paragraphs(tb):
        tb.remove(p)
    for p in new_ps:
        tb.append(p)


# ----------------------------------------------------------------------------- tipi di slot


def fill_static_label(slide: Slide, slot: Slot, ctx: FillContext) -> None:
    sh = shape_by_id(slide, slot.shape_ids[0]) if slot.shape_ids else None
    text = ctx.labels.get(slot.options.get("label_key", ""))
    if sh is None or not text:
        return  # senza traduzione resta il testo originale del template
    _set_single_text(sh, text, ctx)
    if "width_cm" in slot.options:
        sh.width = Cm(slot.options["width_cm"])


def fill_text(slide: Slide, slot: Slot, ctx: FillContext) -> None:
    sh = shape_by_id(slide, slot.shape_ids[0]) if slot.shape_ids else None
    if sh is None:
        return
    value = as_text(field_value(ctx.content, slot.field))
    if not value:
        if slot.field == "full_name":
            _set_single_text(sh, "", ctx)  # nome mancante: lo spazio resta vuoto, da completare a mano
        else:
            remove_shape(sh)
        return
    _set_single_text(sh, value, ctx)
    set_shape_name(sh, _name(slot.id, ctx, chk=f"lines:{int(slot.options.get('max_lines', 1))}"))


def fill_label_lines(slide: Slide, slot: Slot, ctx: FillContext) -> None:
    sh = shape_by_id(slide, slot.shape_ids[0]) if slot.shape_ids else None
    if sh is None:
        return
    ps = paragraphs(txbody(sh))
    if not ps:
        return
    sep = slot.options.get("separator", " ")
    value_bold = slot.options.get("value_bold")
    max_chars = slot.options.get("max_chars")

    new_ps: list[etree._Element] = []
    for i, line in enumerate(slot.options.get("lines", [])):
        value = as_text(field_value(ctx.content, line.get("field")))
        if not value:
            continue
        if max_chars and len(value) > max_chars:
            ctx.warn(slot.id, f"valore troncato a {max_chars} caratteri")
            value = shorten(value, max_chars)
        proto = ps[min(i, len(ps) - 1)]  # ogni riga conserva lo stile della riga corrispondente del template
        base_rpr = first_rpr(proto) if first_rpr(proto) is not None else first_rpr(ps[0])
        label_key = line.get("label_key")
        label = ctx.labels.get(label_key) if label_key else None
        if label_key and not label:  # etichetta non tradotta: si usa il testo della riga del template
            label = "".join(t.text or "" for t in proto.iter(qa("t"))).strip()
        p = copy.deepcopy(proto)
        clear_paragraph_content(p)
        if label:
            label = label.rstrip().rstrip(":") if sep.strip().startswith(":") else label
            append_run(p, make_run(label + sep, style_rpr(base_rpr, lang=ctx.tag)))
            append_run(p, make_run(value, style_rpr(base_rpr, lang=ctx.tag, bold=bool(value_bold))))
        else:
            append_run(p, make_run(value, style_rpr(base_rpr, lang=ctx.tag, bold=True if value_bold else None)))
        _clean_end(p)
        new_ps.append(p)

    if not new_ps:
        remove_shape(sh)
        _remove_ids(slide, slot.remove_shape_ids_if_empty)
        return
    _replace_paragraphs(sh, new_ps)
    set_shape_name(sh, _name(slot.id, ctx, chk=f"lines:{len(new_ps)}"))


def fill_free_text(slide: Slide, slot: Slot, ctx: FillContext) -> None:
    v = field_value(ctx.content, slot.field)
    if isinstance(v, list):
        text = "\n".join(exp_header(e) if isinstance(e, ExperienceBlock) else str(e) for e in v)
    else:
        text = as_text(v)
    if not text:
        return
    sh = _target_shape(slide, slot, ctx)
    if sh is None:
        return
    proto = ctx.protos.get(slot.id, {}).get("para")
    pt = slot.font_pt or para_size_pt(proto) or 12
    space_after = slot.options.get("space_after_pt", 0)
    new_ps = []
    for chunk in text.split("\n"):
        p = paragraph_from_proto(proto if proto is not None else _synth_paragraph(), [(chunk, None)], size_pt=pt, lang=ctx.tag, font=_font(ctx))
        _spacing(p, space_after)
        new_ps.append(p)
    _replace_paragraphs(sh, new_ps)


def fill_bullets(slide: Slide, slot: Slot, ctx: FillContext) -> None:
    v = field_value(ctx.content, slot.field)
    if slot.field == "experiences" and ctx.exp_chunk is not None:
        v = ctx.exp_chunk
    items = as_items(v)
    if not items:
        label = slot.options.get("label_shape_id")  # senza voci non ha senso lasciare il titolo della sezione
        if label is not None:
            _remove_ids(slide, [label])
        sh = shape_by_id(slide, slot.shape_ids[0]) if slot.shape_ids else None
        if sh is not None:
            remove_shape(sh)  # niente segnaposto d'esempio nella slide finale
        return
    sh = _target_shape(slide, slot, ctx)
    if sh is None:
        return
    proto = ctx.protos.get(slot.id, {}).get("para")
    bullet = _with_bullet(proto if proto is not None else _synth_paragraph())
    sub = _with_bullet(bullet, char="–", extra_mar=SUB_STEP_EMU)
    pt = slot.font_pt or para_size_pt(proto) or 11
    space_after = slot.options.get("space_after_pt", 0)
    font = _font(ctx) if proto is None else None
    new_ps = []
    for item in items:
        if isinstance(item, ExperienceBlock):
            role = item.role + (f" ({item.period})" if item.period else "")
            runs = [(item.title, True)] + ([(" – " + role, False)] if role and item.title else [(role, False)] if role else [])
            p = paragraph_from_proto(bullet, runs, size_pt=pt, lang=ctx.tag, font=font)
            _spacing(p, 0)
            new_ps.append(p)
            for i, b in enumerate(item.bullets):
                q = paragraph_from_proto(sub, [(b, False)], size_pt=pt, lang=ctx.tag, font=font)
                _spacing(q, space_after if i == len(item.bullets) - 1 else 0)
                new_ps.append(q)
            if not item.bullets:
                _spacing(p, space_after)
        else:
            p = paragraph_from_proto(bullet, [(str(item), False)], size_pt=pt, lang=ctx.tag, font=font)
            _spacing(p, space_after)
            new_ps.append(p)
    _replace_paragraphs(sh, new_ps)


def fill_pills(slide: Slide, slot: Slot, ctx: FillContext) -> None:
    _remove_ids(slide, slot.remove_shape_ids)
    proto = ctx.protos.get(slot.id, {}).get("pill")
    cols = slot.options.get("columns") or []
    if proto is None or not cols:
        ctx.warn(slot.id, "pillola di riferimento mancante: competenze non inserite")
        return
    skills = [str(s) for s in as_items(field_value(ctx.content, slot.field))]
    rows = int(slot.options.get("max_rows", 1))
    cap = len(cols) * rows
    pt = slot.font_pt or 12
    if len(skills) > cap:
        ctx.warn(slot.id, f"{len(skills) - cap} competenze oltre la capacita' ({cap}) omesse")
        skills = skills[:cap]
    next_id = max_shape_id(slide) + 1
    tree = slide.shapes._spTree
    placed = 0
    for i, skill in enumerate(skills):
        col = cols[placed % len(cols)]
        row = placed // len(cols)
        text_w = ctx.meter.width_pt(skill, pt)
        avail_w = (col["w"] - 2 * PILL_INSET_CM) * CM_TO_PT
        if text_w > avail_w:
            ctx.warn(slot.id, f"competenza '{skill}' troppo lunga per la pillola, omessa")
            continue
        el = copy.deepcopy(proto)
        set_cnv_id(el, next_id, _name(slot.id, ctx, str(i), chk="lines:1"))
        next_id += 1
        xfrm = el.find(f"{{http://schemas.openxmlformats.org/presentationml/2006/main}}spPr").find(qa("xfrm"))
        off, ext = xfrm.find(qa("off")), xfrm.find(qa("ext"))
        off.set("x", str(int(Cm(col["x"]))))
        off.set("y", str(int(Cm(slot.options["first_row_y"] + row * slot.options["row_pitch"]))))
        ext.set("cx", str(int(Cm(col["w"]))))
        ext.set("cy", str(int(Cm(slot.options["pill_h"]))))
        body = txbody(el)
        p0 = paragraphs(body)[0]
        rpr = first_rpr(p0)
        clear_paragraph_content(p0)
        append_run(p0, make_run(skill, style_rpr(rpr, size_pt=pt, lang=ctx.tag)))
        tree.append(el)
        placed += 1


def fill_experience_columns(slide: Slide, slot: Slot, ctx: FillContext) -> None:
    blocks = ctx.exp_chunk if ctx.exp_chunk is not None else ctx.content.experiences
    pt = slot.font_pt or 10
    limit = ctx.spec.bottom_limit_cm
    cols = [sid for sid in slot.shape_ids if shape_by_id(slide, sid) is not None]
    _remove_ids(slide, slot.remove_shape_ids)
    if not cols:
        return
    per_col: list[list[ExperienceBlock]] = [[] for _ in cols]
    for i, blk in enumerate(blocks):
        per_col[i % len(cols)].append(blk)
    pr = ctx.protos[slot.id]
    role_label = ctx.labels.get(slot.options.get("role_label_key", ""), "")

    for ci, sid in enumerate(cols):
        sh = shape_by_id(slide, sid)
        col_blocks = per_col[ci]
        if not col_blocks:
            remove_shape(sh)
            continue
        top_cm = Emu(sh.top).cm
        sh.height = Cm(limit - top_cm)
        if slot.options.get("column_width_cm"):
            sh.width = Cm(slot.options["column_width_cm"])
        tb = txbody(sh)
        for p in paragraphs(tb):
            tb.remove(p)
        for bi, blk in enumerate(col_blocks):
            if bi > 0:
                tb.append(paragraph_from_proto(pr["blank"], [(" ", None)], size_pt=pt, lang=ctx.tag))
            tb.append(paragraph_from_proto(pr["title"], [(blk.title, True)], size_pt=pt, lang=ctx.tag))
            role_text = blk.role + (f" ({blk.period})" if blk.period else "")
            role_runs = ([(role_label + " ", False)] if role_label else []) + [(role_text, False)]
            tb.append(paragraph_from_proto(pr["role"], role_runs, size_pt=pt, lang=ctx.tag))
            tb.append(paragraph_from_proto(pr["blank"], [(" ", None)], size_pt=pt, lang=ctx.tag))
            for bullet in blk.bullets:
                tb.append(paragraph_from_proto(pr["bullet"], [(bullet, False)], size_pt=pt, lang=ctx.tag))
        set_shape_name(sh, _name(slot.id, ctx, f"c{ci}", chk="h"))


KIND_HANDLERS = {
    "static_label": fill_static_label,
    "text": fill_text,
    "label_lines": fill_label_lines,
    "free_text": fill_free_text,
    "bullets": fill_bullets,
    "pills": fill_pills,
    "experience_columns": fill_experience_columns,
}


def fill_slide(slide: Slide, slide_key: str, ctx: FillContext) -> None:
    for slot in ctx.spec.slots_for(slide_key):
        if slot.kind not in ("pills", "experience_columns"):
            _remove_ids(slide, [i for i in slot.remove_shape_ids if i not in slot.shape_ids])
        KIND_HANDLERS[slot.kind](slide, slot, ctx)


def slot_areas(prs: Presentation, spec: TemplateSpec) -> dict[str, Box | None]:
    return {s.id: slot_area(prs, spec, s) for s in spec.slots}


_ = A_NS  # re-export
