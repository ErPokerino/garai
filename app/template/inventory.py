"""Inventario delle shape di un template PPTX (per validazione umana o proposta di Template Spec via LLM)."""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

from .xmlutil import iter_shapes


def inventory(pptx_path: str | Path) -> dict:
    prs = Presentation(str(pptx_path))
    slides = []
    for i, slide in enumerate(prs.slides):
        shapes = []
        for sh in iter_shapes(slide.shapes):
            entry = {
                "id": sh.shape_id,
                "name": sh.name,
                "type": str(sh.shape_type),
                "x_cm": round(Emu(sh.left or 0).cm, 2),
                "y_cm": round(Emu(sh.top or 0).cm, 2),
                "w_cm": round(Emu(sh.width or 0).cm, 2),
                "h_cm": round(Emu(sh.height or 0).cm, 2),
            }
            if sh.has_text_frame:
                entry["paragraphs"] = [
                    {
                        "text": p.text,
                        "level": p.level,
                        "size_pt": next((r.font.size.pt for r in p.runs if r.font.size), None),
                        "bold": next((r.font.bold for r in p.runs if r.font.bold is not None), None),
                    }
                    for p in sh.text_frame.paragraphs
                ]
            shapes.append(entry)
        slides.append({"index": i, "layout": slide.slide_layout.name, "shapes": shapes})
    return {
        "slide_w_cm": round(Emu(prs.slide_width).cm, 2),
        "slide_h_cm": round(Emu(prs.slide_height).cm, 2),
        "slides": slides,
    }
