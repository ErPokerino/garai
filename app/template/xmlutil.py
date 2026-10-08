"""Helper di basso livello per manipolare le shape e i paragrafi DrawingML preservando gli stili."""
from __future__ import annotations

import copy
from typing import Iterator

from lxml import etree
from pptx.shapes.base import BaseShape
from pptx.shapes.group import GroupShape
from pptx.slide import Slide

A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NS = {"a": A_NS, "p": P_NS, "r": R_NS}


def qa(tag: str) -> str:
    return f"{{{A_NS}}}{tag}"


def iter_shapes(shapes) -> Iterator[BaseShape]:
    for sh in shapes:
        yield sh
        if isinstance(sh, GroupShape):
            yield from iter_shapes(sh.shapes)


def shape_by_id(slide: Slide, shape_id: int) -> BaseShape | None:
    for sh in iter_shapes(slide.shapes):
        if sh.shape_id == shape_id:
            return sh
    return None


def remove_shape(shape: BaseShape) -> None:
    el = shape._element
    el.getparent().remove(el)


def max_shape_id(slide: Slide) -> int:
    ids = [int(x) for x in slide.shapes._spTree.xpath(".//p:cNvPr/@id")]
    return max(ids) if ids else 1


def txbody(shape_or_el) -> etree._Element:
    el = getattr(shape_or_el, "_element", shape_or_el)
    tb = el.find(f"{{{P_NS}}}txBody")
    if tb is None:
        raise ValueError("shape senza txBody")
    return tb


def paragraphs(tb: etree._Element) -> list[etree._Element]:
    return tb.findall(qa("p"))


def clear_paragraph_content(p: etree._Element) -> None:
    for child in list(p):
        if child.tag in (qa("r"), qa("br"), qa("fld")):
            p.remove(child)


def first_rpr(p: etree._Element) -> etree._Element | None:
    r = p.find(qa("r"))
    if r is not None:
        rpr = r.find(qa("rPr"))
        if rpr is not None:
            return copy.deepcopy(rpr)
    return None


def style_rpr(
    rpr: etree._Element | None,
    *,
    size_pt: float | None = None,
    bold: bool | None = None,
    lang: str | None = None,
    font: str | None = None,
) -> etree._Element:
    rpr = copy.deepcopy(rpr) if rpr is not None else etree.Element(qa("rPr"))
    for attr in ("err", "dirty", "altLang"):
        if attr in rpr.attrib:
            del rpr.attrib[attr]
    if size_pt is not None:
        rpr.set("sz", str(int(round(size_pt * 100))))
    if bold is not None:
        rpr.set("b", "1" if bold else "0")
    if lang:
        rpr.set("lang", lang)
    if font and rpr.find(qa("latin")) is None:
        latin = etree.SubElement(rpr, qa("latin"))
        latin.set("typeface", font)
    return rpr


def make_run(text: str, rpr: etree._Element | None) -> etree._Element:
    r = etree.Element(qa("r"))
    if rpr is not None:
        r.append(copy.deepcopy(rpr))
    t = etree.SubElement(r, qa("t"))
    t.text = text
    return r


def append_run(p: etree._Element, run: etree._Element) -> None:
    end = p.find(qa("endParaRPr"))
    if end is not None:
        end.addprevious(run)
    else:
        p.append(run)


def paragraph_from_proto(
    proto: etree._Element,
    runs: list[tuple[str, bool | None]],
    *,
    size_pt: float | None,
    lang: str,
    font: str | None = None,
    rpr_source: etree._Element | None = None,
) -> etree._Element:
    """Clona un paragrafo del template e ne sostituisce i run mantenendo pPr e stile del primo run."""
    p = copy.deepcopy(proto)
    base = rpr_source if rpr_source is not None else first_rpr(proto)
    clear_paragraph_content(p)
    for text, bold in runs:
        append_run(p, make_run(text, style_rpr(base, size_pt=size_pt, bold=bold, lang=lang, font=font)))
    end = p.find(qa("endParaRPr"))
    if end is not None and "err" in end.attrib:
        del end.attrib["err"]
    return p


def _cnv_pr(el: etree._Element) -> etree._Element:
    return next(el.iter(f"{{{P_NS}}}cNvPr"))


def set_shape_name(shape: BaseShape, name: str) -> None:
    _cnv_pr(shape._element).set("name", name)


def set_cnv_id(el: etree._Element, new_id: int, name: str | None = None) -> None:
    c = _cnv_pr(el)
    c.set("id", str(new_id))
    if name:
        c.set("name", name)
    # il creationId duplicato non e' un problema ma lo rimuoviamo per evitare conflitti
    for ext in c.findall(qa("extLst")):
        c.remove(ext)
