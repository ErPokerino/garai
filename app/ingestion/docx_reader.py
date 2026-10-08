"""Lettura di DOCX in testo strutturato (markdown-like) con ordine documento, tabelle, header/footer e caselle di testo."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph


def _heading_prefix(style_name: str) -> str:
    s = (style_name or "").lower()
    if s.startswith("heading") or s.startswith("titolo"):
        digits = "".join(ch for ch in s if ch.isdigit())
        level = int(digits) if digits else 1
        return "#" * min(max(level, 1), 6) + " "
    if s == "title":
        return "# "
    return ""


def _is_list(p: Paragraph) -> bool:
    ppr = p._p.pPr
    return ppr is not None and ppr.numPr is not None


def _para_text(p: Paragraph) -> str:
    # include anche i testi in w:hyperlink (python-docx li espone in p.text nelle versioni recenti)
    return p.text.strip()


def _iter_block_items(parent):
    """Itera paragrafi e tabelle in ordine documento (parent: Document o _Cell)."""
    container = parent._tc if isinstance(parent, _Cell) else parent.element.body
    for child in container.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, parent)
        elif child.tag == qn("w:tbl"):
            yield Table(child, parent)


def _table_lines(table: Table, depth: int = 0) -> list[str]:
    lines: list[str] = []
    for row in table.rows:
        cells: list[str] = []
        seen = set()
        for c in row.cells:
            if id(c._tc) in seen:
                continue
            seen.add(id(c._tc))
            cells.append(_cell_text(c, depth))
        if any(cells):
            lines.append(" | ".join(cells))
    return lines


def _cell_text(cell: _Cell, depth: int = 0) -> str:
    parts: list[str] = []
    for block in _iter_block_items(cell):
        if isinstance(block, Paragraph):
            t = _para_text(block)
            if t:
                parts.append(("- " if _is_list(block) else "") + t)
        elif depth < 3:
            parts.extend(_table_lines(block, depth + 1))
    return " / ".join(parts)


def _textboxes(root_element) -> list[str]:
    out: list[str] = []
    seen: set[str] = set()
    for tx in root_element.iter(qn("w:txbxContent")):
        txt = " ".join("".join(t.text or "" for t in p.iter(qn("w:t"))).strip() for p in tx.iter(qn("w:p"))).strip()
        if txt and txt not in seen:  # AlternateContent duplica Choice/Fallback
            seen.add(txt)
            out.append(txt)
    return out


def read_docx(path: str | Path) -> str:
    doc = Document(str(path))
    lines: list[str] = []

    # header/footer (spesso contengono nome, codice, contatti)
    hf: list[str] = []
    for section in doc.sections:
        for part in (section.header, section.footer, section.first_page_header, section.first_page_footer):
            try:
                for p in part.paragraphs:
                    t = p.text.strip()
                    if t:
                        hf.append(t)
                for tb in part.tables:
                    hf.extend(_table_lines(tb))
                hf.extend(_textboxes(part._element))
            except Exception:
                continue
    hf = list(dict.fromkeys(hf))
    if hf:
        lines.append("[HEADER/FOOTER] " + " | ".join(hf))

    for block in _iter_block_items(doc):
        if isinstance(block, Paragraph):
            t = _para_text(block)
            if not t:
                continue
            prefix = _heading_prefix(block.style.name if block.style is not None else "")
            if not prefix and _is_list(block):
                prefix = "- "
            lines.append(prefix + t)
        else:
            lines.extend(_table_lines(block))

    boxes = _textboxes(doc.element.body)
    if boxes:
        lines.append("[TEXTBOX] " + " | ".join(boxes))
    return "\n".join(lines)
