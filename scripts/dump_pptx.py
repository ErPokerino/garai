"""Dump delle shape di un PPTX (nome, tipo, posizione, testo, font) per ispezione del template."""
import sys

from pptx import Presentation
from pptx.util import Emu


def walk(shapes, depth=0):
    for sh in shapes:
        pad = "  " * depth
        pos = f"x={Emu(sh.left or 0).cm:.2f} y={Emu(sh.top or 0).cm:.2f} w={Emu(sh.width or 0).cm:.2f} h={Emu(sh.height or 0).cm:.2f}"
        print(f"{pad}- id={sh.shape_id} name={sh.name!r} type={sh.shape_type} {pos}")
        if sh.is_placeholder:
            print(f"{pad}  placeholder idx={sh.placeholder_format.idx} type={sh.placeholder_format.type}")
        if sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                runs = [(r.text, r.font.size.pt if r.font.size else None, r.font.name, r.font.bold) for r in p.runs]
                if runs:
                    print(f"{pad}  P(lvl={p.level}): {runs}")
        if getattr(sh, "has_table", False) and sh.has_table:
            for r in sh.table.rows:
                print(f"{pad}  ROW: " + " | ".join(c.text.replace("\n", " / ") for c in r.cells))
        if sh.shape_type == 6:  # group
            walk(sh.shapes, depth + 1)


def main(path):
    prs = Presentation(path)
    print(f"slide size: {Emu(prs.slide_width).cm:.2f} x {Emu(prs.slide_height).cm:.2f} cm")
    print("layouts:", [l.name for l in prs.slide_layouts])
    for i, s in enumerate(prs.slides, 1):
        print(f"=== slide {i} layout={s.slide_layout.name!r}")
        walk(s.shapes)


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main(sys.argv[1])
